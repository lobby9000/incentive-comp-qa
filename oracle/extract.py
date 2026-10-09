"""
Read the recalculated oracle and write its results as plain files for the
reconciliation step (the API side never opens the workbook).

Outputs:
  oracle/oracle_payouts.csv   one row per (period, seller)
  oracle/oracle_summary.json  period totals, control checks and dataset hash

Usage: python oracle/extract.py
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
GEN = HERE.parent / "data" / "generated"


def cents(value) -> int:
    return int(round(float(value) * 100))


def main() -> None:
    wb = load_workbook(HERE / "commission_oracle.xlsx", data_only=True, read_only=True)
    rows = []
    for r in wb["Calc"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        period, seller, quota, quota_c, credited_c, _cr, hundredths, _pct, qualified, *_, payout_c, _payout = r
        if any(v is None for v in (credited_c, hundredths, payout_c)):
            raise SystemExit("The oracle has no cached values. Run oracle/recalc.py first.")
        rows.append({
            "period_id": period,
            "seller_id": seller,
            "quota": f"{int(quota_c) / 100:.2f}",
            "net_credited": f"{int(credited_c) / 100:.2f}",
            "attainment_pct": f"{int(hundredths) / 100:.2f}",
            "qualified": int(qualified) == 1,
            "payout": f"{int(payout_c) / 100:.2f}",
        })
    with open(HERE / "oracle_payouts.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

    checks = [r[1] for r in wb["Checks"].iter_rows(min_row=4, values_only=True) if r[0]]
    summary = {}
    for r in wb["Summary"].iter_rows(min_row=4, max_row=6, values_only=True):
        summary[r[0]] = {"payout_rows": int(r[1]), "qualified": int(r[2]),
                         "total_quota": f"{r[3]:.2f}", "net_credited": f"{r[4]:.2f}",
                         "total_payout": f"{r[5]:.2f}"}
    manifest = json.loads((GEN / "MANIFEST.json").read_text(encoding="utf-8"))
    out = {
        "dataset_sha256": manifest["sha256"]["transactions.csv"],
        "control_checks_all_zero": all(int(c) == 0 for c in checks),
        "periods": summary,
    }
    (HERE / "oracle_summary.json").write_text(json.dumps(out, indent=2) + "\n", encoding="utf-8")
    print(f"{len(rows)} payout rows written; control checks zero: {out['control_checks_all_zero']}")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
