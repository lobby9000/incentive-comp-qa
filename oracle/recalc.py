"""
Recalculate every formula in an .xlsx with headless LibreOffice and report errors.

openpyxl writes formulas without cached values; this opens the workbook in
LibreOffice, runs calculateAll(), saves it in place and then scans every cell
for spreadsheet errors (#DIV/0!, #N/A, #VALUE!, ...).

Usage: python oracle/recalc.py oracle/commission_oracle.xlsx [timeout_seconds]
Exit code 0 only when the file was recalculated and contains no error values.
Requires LibreOffice (`soffice`) on PATH; on Ubuntu: apt-get install libreoffice-calc
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""

ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def recalc(path: Path, timeout: int) -> dict:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        return {"error": "LibreOffice (soffice) not found on PATH"}
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
        url = Path(profile).as_uri()
        subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=timeout, env=env)
        macro_dir = Path(profile) / "user" / "basic" / "Standard"
        if not macro_dir.exists():
            return {"error": "LibreOffice did not create a user profile"}
        (macro_dir / "Module1.xba").write_text(MACRO)
        before = path.stat().st_mtime_ns
        result = subprocess.run(
            [soffice, "--headless", "--norestore", f"-env:UserInstallation={url}",
             "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
             str(path.resolve())],
            capture_output=True, text=True, timeout=timeout, env=env)
        if result.returncode != 0:
            return {"error": f"soffice exited {result.returncode}: {result.stderr.strip()}"}
        if path.stat().st_mtime_ns == before:
            return {"error": "LibreOffice did not rewrite the file; nothing was recalculated"}

    wb = load_workbook(path, data_only=True)
    found: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERRORS:
                    found.setdefault(cell.value, []).append(f"{ws.title}!{cell.coordinate}")
    total = sum(len(v) for v in found.values())
    return {"status": "success" if total == 0 else "errors_found", "total_errors": total,
            "errors": {k: v[:20] for k, v in found.items()}}


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = Path(sys.argv[1])
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 600
    report = recalc(path, timeout)
    print(json.dumps(report, indent=2))
    return 0 if report.get("status") == "success" else 1


if __name__ == "__main__":
    sys.exit(main())
