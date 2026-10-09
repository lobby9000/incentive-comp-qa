"""
Build the independent Excel oracle for plan RET-Q4-2026.

The workbook re-implements requirements 3.1 to 3.3 with spreadsheet formulas,
independently of the Python engine, so its payouts can be reconciled against
the API (a parallel run). Every number in the Calc sheet is a live formula over
the raw transactions; nothing is pasted from Python.

Money is handled in integer units to avoid binary floating point:
  * amounts are converted to cents,
  * band limits are compared in "cent-percent" units (cents x 100),
  * band commissions are accumulated as cent-percent x basis points, an exact
    integer, and divided by 1,000,000 only once, inside the final half-up round.
A plain =ROUND(SUM(...),2) over decimal floats mis-rounds some half-cent ties,
which is exactly the case requirement REQ-CAL-06 pins down.

Usage:
    python oracle/build_oracle.py            # writes oracle/commission_oracle.xlsx
    python oracle/recalc.py oracle/commission_oracle.xlsx
    python oracle/extract.py                 # writes oracle/oracle_payouts.csv
"""

from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
GEN = ROOT / "data" / "generated"
OUT = Path(__file__).resolve().parent / "commission_oracle.xlsx"

ARIAL = "Arial"
F_INPUT = Font(name=ARIAL, size=10, color="0000FF")
F_FORMULA = Font(name=ARIAL, size=10, color="000000")
F_LINK = Font(name=ARIAL, size=10, color="008000")
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_BODY = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")


def header(ws, row, labels, widths=None):
    for j, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=label)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w


def build(gen: Path = GEN, out: Path = OUT):
    manifest = json.loads((gen / "MANIFEST.json").read_text(encoding="utf-8"))
    txns = list(csv.DictReader(open(gen / "transactions.csv", encoding="utf-8")))
    quotas = list(csv.DictReader(open(gen / "quotas.csv", encoding="utf-8")))
    n = len(txns)
    last = n + 1  # last data row in Txns
    periods = ["2026-10", "2026-11", "2026-12"]

    wb = Workbook()

    # ------------------------------------------------------------------ README
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Commission oracle for plan RET-Q4-2026", F_TITLE),
        ("", F_BODY),
        ("Purpose: an independent, formula-only re-implementation of requirements 3.1 to 3.3, used to "
         "reconcile every payout produced by the Commission Engine API (parallel run).", F_BODY),
        ("Inputs (blue): the raw dataset in Txns columns A:J and Calc columns A:C, and the plan "
         "parameters in Plan. Formulas (black) compute everything else. Green = link to another sheet.", F_BODY),
        ("Integer arithmetic: cents for money, cent-percent (cents x 100) for band limits, and basis "
         "points for rates. The exact total is divided by 1,000,000 only inside the final half-up "
         "round, so half-cent ties (REQ-CAL-06) round the same way as exact decimal arithmetic.", F_BODY),
        ("Txns: one row per transaction. Returns look up their original sale (MATCH/INDEX) to inherit "
         "sellers and splits and to measure days since the sale for the 30-day clawback window.", F_BODY),
        ("Calc: one row per quota (period, seller), which is also one payout row per REQ-CLS-03.", F_BODY),
        ("Checks: control totals that must be zero. Summary: period totals for the test summary report.", F_BODY),
        ("", F_BODY),
        (f"Dataset: seed {manifest['seed']}, {manifest['sales']:,} sales, {manifest['returns']:,} returns, "
         f"{manifest['sellers']} sellers, {manifest['quotas']} quotas.", F_BODY),
        (f"transactions.csv SHA-256: {manifest['sha256']['transactions.csv']}", F_BODY),
        ("Rebuild after regenerating data: python oracle/build_oracle.py, then python oracle/recalc.py "
         "oracle/commission_oracle.xlsx, then python oracle/extract.py.", F_BODY),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ------------------------------------------------------------------ Plan
    ws = wb.create_sheet("Plan")
    ws.column_dimensions["A"].width = 34
    for col in "BCDE":
        ws.column_dimensions[col].width = 16
    ws["A1"] = "Plan parameters (RET-Q4-2026, requirements 3.2 and 3.3)"
    ws["A1"].font = F_TITLE
    params = [
        (4, "Gate, % of quota (REQ-CAL-03)", 50),
        (5, "Clawback window, days (REQ-RET-04)", 30),
    ]
    for r, label, value in params:
        ws.cell(row=r, column=1, value=label).font = F_BODY
        c = ws.cell(row=r, column=2, value=value)
        c.font = F_INPUT
        c.fill = FILL_KEY
    header(ws, 8, ["Band (REQ-CAL-04)", "Lower % of quota", "Upper % of quota", "Rate (bp)", "Rate %"])
    bands = [("A", 0, 80, 100), ("B", 80, 100, 250), ("C", 100, 150, 400), ("D", 150, 200, 500)]
    for i, (name, lo, hi, bp) in enumerate(bands, start=9):
        ws.cell(row=i, column=1, value=name).font = F_BODY
        for j, v in enumerate((lo, hi, bp), start=2):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_INPUT
        c = ws.cell(row=i, column=5, value=f"=D{i}/100")
        c.font = F_FORMULA
        c.number_format = "0.00"
    ws["A14"] = "Above the upper limit of band D the rate is 0% (cap). Source: requirements v1.0, section 3.3."
    ws["A14"].font = F_BODY

    # ------------------------------------------------------------------ Txns
    ws = wb.create_sheet("Txns")
    cols = ["transaction_id", "type", "txn_date", "store_id", "amount", "seller_1", "split_1",
            "seller_2", "split_2", "original_transaction_id",
            "amount_cents", "period", "orig_row", "days_since_sale", "eligible",
            "eff_seller_1", "eff_seller_2", "eff_split_2", "sign", "credit_2_cents", "credit_1_cents"]
    header(ws, 1, cols, [16, 8, 11, 9, 12, 9, 7, 9, 7, 16, 12, 9, 9, 9, 8, 10, 10, 9, 6, 12, 12])
    ws.freeze_panes = "B2"
    rng = lambda col: f"${col}$2:${col}${last}"  # noqa: E731
    for i, t in enumerate(txns, start=2):
        values = [
            t["transaction_id"], t["type"], date.fromisoformat(t["txn_date"]), t["store_id"],
            float(t["amount"]), t["seller_1"] or None, int(t["split_1"]) if t["split_1"] else None,
            t["seller_2"] or None, int(t["split_2"]) if t["split_2"] else None,
            t["original_transaction_id"] or None,
        ]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_INPUT
        ws.cell(row=i, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=5).number_format = "#,##0.00"
        formulas = [
            f"=ROUND(E{i}*100,0)",
            f'=YEAR(C{i})&"-"&TEXT(MONTH(C{i}),"00")',
            f'=IF(B{i}="RETURN",MATCH(J{i},{rng("A")},0),"")',
            f'=IF(B{i}="RETURN",C{i}-INDEX({rng("C")},M{i}),"")',
            f'=IF(B{i}="SALE",1,IF(N{i}<=Plan!$B$5,1,0))',
            f'=IF(B{i}="SALE",F{i},INDEX({rng("F")},M{i}))',
            f'=IF(B{i}="SALE",IF(H{i}="","",H{i}),IF(INDEX({rng("H")},M{i})="","",INDEX({rng("H")},M{i})))',
            f'=IF(Q{i}="",0,IF(B{i}="SALE",I{i},INDEX({rng("I")},M{i})))',
            f'=IF(B{i}="SALE",1,-1)',
            f'=IF(Q{i}="",0,S{i}*O{i}*INT(K{i}*R{i}/100))',
            f"=S{i}*O{i}*K{i}-T{i}",
        ]
        for j, f in enumerate(formulas, start=11):
            c = ws.cell(row=i, column=j, value=f)
            c.font = F_FORMULA

    # ------------------------------------------------------------------ Calc
    ws = wb.create_sheet("Calc")
    cols = ["period", "seller_id", "quota", "quota_cents", "credited_cents", "credited",
            "attainment_hundredths", "attainment_pct", "qualified",
            "base_A_cp", "base_B_cp", "base_C_cp", "base_D_cp", "total_scaled", "payout_cents", "payout"]
    header(ws, 1, cols, [9, 10, 13, 13, 14, 13, 12, 12, 9, 14, 14, 14, 14, 18, 12, 12])
    ws.freeze_panes = "C2"
    q_last = len(quotas) + 1
    for i, q in enumerate(quotas, start=2):
        for j, v in enumerate((q["period_id"], q["seller_id"], float(q["amount"])), start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = F_INPUT
        ws.cell(row=i, column=3).number_format = "#,##0.00"
        band_rows = (9, 10, 11, 12)
        formulas = [
            f"=ROUND(C{i}*100,0)",
            (f"=SUMIFS(Txns!$U$2:$U${last},Txns!$P$2:$P${last},B{i},Txns!$L$2:$L${last},A{i})"
             f"+SUMIFS(Txns!$T$2:$T${last},Txns!$Q$2:$Q${last},B{i},Txns!$L$2:$L${last},A{i})"),
            f"=E{i}/100",
            f"=SIGN(E{i})*INT((ABS(E{i})*20000+D{i})/(2*D{i}))",
            f"=G{i}/100",
            f"=IF(E{i}*100>=D{i}*Plan!$B$4,1,0)",
        ] + [
            f"=MAX(0,MIN(E{i}*100,D{i}*Plan!$C${r})-D{i}*Plan!$B${r})" for r in band_rows
        ] + [
            f"=J{i}*Plan!$D$9+K{i}*Plan!$D$10+L{i}*Plan!$D$11+M{i}*Plan!$D$12",
            f"=IF(I{i}=1,INT((N{i}+500000)/1000000),0)",
            f"=O{i}/100",
        ]
        for j, f in enumerate(formulas, start=4):
            c = ws.cell(row=i, column=j, value=f)
            c.font = F_LINK if j == 5 else F_FORMULA
        for col in (6, 16):
            ws.cell(row=i, column=col).number_format = "#,##0.00"
        ws.cell(row=i, column=8).number_format = "0.00"

    # ------------------------------------------------------------------ Checks
    ws = wb.create_sheet("Checks")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 16
    ws["A1"] = "Control checks (every value must be 0)"
    ws["A1"].font = F_TITLE
    header(ws, 3, ["Check", "Value"])
    r = 4
    for p in periods:
        ws.cell(row=r, column=1, value=f"{p}: credit cents not covered by a quota row (REQ-CLS-02)").font = F_BODY
        c = ws.cell(row=r, column=2, value=(
            f"=SUMIFS(Txns!$U$2:$U${last},Txns!$L$2:$L${last},\"{p}\")"
            f"+SUMIFS(Txns!$T$2:$T${last},Txns!$L$2:$L${last},\"{p}\")"
            f"-SUMIFS(Calc!$E$2:$E${q_last},Calc!$A$2:$A${q_last},\"{p}\")"))
        c.font = F_LINK
        r += 1
    ws.cell(row=r, column=1, value="Sale credits minus sale amounts, in cents (REQ-CRD-02)").font = F_BODY
    c = ws.cell(row=r, column=2, value=(
        f"=SUMIFS(Txns!$U$2:$U${last},Txns!$B$2:$B${last},\"SALE\")"
        f"+SUMIFS(Txns!$T$2:$T${last},Txns!$B$2:$B${last},\"SALE\")"
        f"-SUMIFS(Txns!$K$2:$K${last},Txns!$B$2:$B${last},\"SALE\")"))
    c.font = F_LINK
    r += 1
    ws.cell(row=r, column=1, value="Returns that do not find their original sale").font = F_BODY
    c = ws.cell(row=r, column=2, value=f'=COUNTIFS(Txns!$B$2:$B${last},"RETURN")-COUNT(Txns!$M$2:$M${last})')
    c.font = F_LINK

    # ------------------------------------------------------------------ Summary
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Period summary (oracle)"
    ws["A1"].font = F_TITLE
    header(ws, 3, ["Period", "Payout rows", "Qualified", "Total quota", "Net credited", "Total payout"],
           [10, 12, 11, 18, 18, 16])
    for k, p in enumerate(periods, start=4):
        ws.cell(row=k, column=1, value=p).font = F_INPUT
        cells = [
            f'=COUNTIFS(Calc!$A$2:$A${q_last},A{k})',
            f'=COUNTIFS(Calc!$A$2:$A${q_last},A{k},Calc!$I$2:$I${q_last},1)',
            f'=SUMIFS(Calc!$C$2:$C${q_last},Calc!$A$2:$A${q_last},A{k})',
            f'=SUMIFS(Calc!$F$2:$F${q_last},Calc!$A$2:$A${q_last},A{k})',
            f'=SUMIFS(Calc!$P$2:$P${q_last},Calc!$A$2:$A${q_last},A{k})',
        ]
        for j, f in enumerate(cells, start=2):
            c = ws.cell(row=k, column=j, value=f)
            c.font = F_LINK
            if j >= 4:
                c.number_format = "#,##0.00"
    ws.cell(row=7, column=1, value="Total").font = F_BOLD
    for j, col in enumerate("BCDEF", start=2):
        c = ws.cell(row=7, column=j, value=f"=SUM({col}4:{col}6)")
        c.font = F_BOLD
        if j >= 4:
            c.number_format = "#,##0.00"

    wb.save(out)
    print(f"wrote {out} ({n:,} transactions, {len(quotas)} quota rows)")


if __name__ == "__main__":
    import sys

    if len(sys.argv) == 3:  # python oracle/build_oracle.py <generated_dir> <out.xlsx>
        build(Path(sys.argv[1]), Path(sys.argv[2]))
    else:
        build()
